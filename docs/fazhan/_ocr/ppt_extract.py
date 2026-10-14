"""课件 (PPT) → courseware/<book>/lNN.md. .ppt → soffice headless → .pptx → python-pptx
текст (слайд бүрээр, хүснэгт орно). Зураг/зургийн слайд (牡丹花 г.м.) зөвхөн гарчиг үлдэнэ.
Хэрэглээ: python3 ppt_extract.py <book> <src_dir> <out_dir>"""
import os, re, sys, glob, subprocess, tempfile
from pptx import Presentation
book, src_dir, out_dir = sys.argv[1], sys.argv[2], sys.argv[3]
os.makedirs(out_dir, exist_ok=True)
tmp = tempfile.mkdtemp()
files = sorted(glob.glob(os.path.join(src_dir, "*.ppt*")), key=lambda f: int(re.search(r"第(\d+)课", f).group(1)))
for src in files:
    k = int(re.search(r"第(\d+)课", src).group(1)); out = os.path.join(out_dir, f"l{k:02d}.md")
    if os.path.exists(out): continue
    f = src
    if src.lower().endswith(".ppt"):
        subprocess.run(["soffice", "--headless", "--convert-to", "pptx", "--outdir", tmp, src], capture_output=True, timeout=300)
        f = os.path.join(tmp, os.path.splitext(os.path.basename(src))[0] + ".pptx")
    prs = Presentation(f); md = [f"# {book} · 第{k}课 课件 (PPT)", "", f"_{len(prs.slides)} слайд. Эх: `{os.path.basename(src)}`_", ""]
    for i, s in enumerate(prs.slides, 1):
        parts = []
        for sh in s.shapes:
            if sh.has_text_frame:
                t = "\n".join(p.text for p in sh.text_frame.paragraphs if p.text.strip())
                if t.strip(): parts.append(t)
            if getattr(sh, "has_table", False) and sh.has_table:
                parts += [" | ".join(c.text.strip() for c in r.cells) for r in sh.table.rows]
        md += [f"## Слайд {i}", "\n".join(parts) or "_(зураг)_", ""]
    open(out, "w", encoding="utf-8").write("\n".join(md)); print(f"L{k}: {len(prs.slides)} slides")
